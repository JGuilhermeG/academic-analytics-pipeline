import os
import sqlite3
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE_DIR, 'academic_tracker.db')
CAMINHO_EXCEL = os.path.join(BASE_DIR, "Relatorio_Academico.xlsx")

QUERY = '''
WITH notas_pivot AS (
    SELECT 
        m.cod_disc, m.disciplina, 6.0 as meta_aprovacao, 
        MAX(CASE WHEN a.prova = 'P1' THEN a.nota_obtida END) as P1,
        MAX(CASE WHEN a.prova = 'P2' THEN a.nota_obtida END) as P2,
        MAX(CASE WHEN a.prova = 'P3' THEN a.nota_obtida END) as P3,
        MAX(CASE WHEN a.prova = 'P4' THEN a.nota_obtida END) as P4,
        MAX(CASE WHEN a.prova = 'SUB1' THEN a.nota_obtida END) as SUB1,
        MAX(CASE WHEN a.prova = 'SUB2' THEN a.nota_obtida END) as SUB2,
        MAX(CASE WHEN a.prova = 'T1' THEN a.nota_obtida END) as T1,
        MAX(CASE WHEN a.prova = 'T2' THEN a.nota_obtida END) as T2,
        MAX(CASE WHEN a.prova = 'T3' THEN a.nota_obtida END) as T3,
        MAX(CASE WHEN a.prova = 'T4' THEN a.nota_obtida END) as T4
    FROM materias m
    LEFT JOIN avaliacoes a ON m.id_turma_disc = a.id_turma_disc
    WHERE m.cod_disc IS NOT NULL
    GROUP BY m.cod_disc, m.disciplina
),
notas_substituidas AS (
    SELECT
        cod_disc, disciplina, meta_aprovacao,
        MAX(COALESCE(P1, 0), COALESCE(SUB1, 0)) as P1_final,
        MAX(COALESCE(P2, 0), COALESCE(SUB1, 0)) as P2_final,
        MAX(COALESCE(P3, 0), COALESCE(SUB2, 0)) as P3_final,
        MAX(COALESCE(P4, 0), COALESCE(SUB2, 0)) as P4_final,
        COALESCE(T1, 0) as T1_final,
        COALESCE(T2, 0) as T2_final,
        COALESCE(T3, 0) as T3_final,
        COALESCE(T4, 0) as T4_final
    FROM notas_pivot
),
notas_calculadas AS (
    SELECT 
        ns.cod_disc, 
        ns.disciplina,
        ns.P1_final, ns.P2_final, ns.P3_final, ns.P4_final,
        ns.T1_final, ns.T2_final, ns.T3_final, ns.T4_final,
        ns.meta_aprovacao,
        ROUND(
            (ns.P1_final * COALESCE(p.peso_p1, 0)) + 
            (ns.P2_final * COALESCE(p.peso_p2, 0)) + 
            (ns.P3_final * COALESCE(p.peso_p3, 0)) + 
            (ns.P4_final * COALESCE(p.peso_p4, 0)) + 
            (ns.T1_final * COALESCE(p.peso_t1, 0)) + 
            (ns.T2_final * COALESCE(p.peso_t2, 0)) + 
            (ns.T3_final * COALESCE(p.peso_t3, 0)) + 
            (ns.T4_final * COALESCE(p.peso_t4, 0))
        , 2) as nota_atual
    FROM notas_substituidas ns
    LEFT JOIN pesos p ON ns.cod_disc = p.cod_disc
)
SELECT 
    cod_disc as "Código", 
    disciplina as "Disciplina", 
    P1_final as "P1", P2_final as "P2", P3_final as "P3", P4_final as "P4",
    T1_final as "T1", T2_final as "T2", T3_final as "T3", T4_final as "T4",
    meta_aprovacao as "Meta",
    nota_atual as "Nota Atual",
    CASE 
        WHEN nota_atual >= meta_aprovacao THEN 'Aprovado'
        ELSE 'Faltam ' || ROUND(meta_aprovacao - nota_atual, 2) || ' pts'
    END as "Status"
FROM notas_calculadas;
'''

def gerar_relatorio():
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query(QUERY, conn)
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados Acadêmicos"

    headers = list(df.columns)
    ws.append(headers)

    for row in df.values:
        ws.append(list(row))

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=len(df) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border_thin
            if cell.column in [1, 2]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    ws.views.sheetView[0].showGridLines = True
    wb.save(CAMINHO_EXCEL)
    print(f"[SUCCESS] Relatório exportado com sucesso: {CAMINHO_EXCEL}")

if __name__ == '__main__':
    gerar_relatorio()